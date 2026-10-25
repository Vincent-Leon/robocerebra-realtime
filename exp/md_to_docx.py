"""把本项目报告用到的 Markdown 子集转成 Word（.docx），用于导入飞书文档。
支持：# / ## / ### 标题、段落（**加粗**、`代码`）、- 无序列表（两级，缩进 2 空格为第二级）、1. 有序列表（每个列表重新编号）、
列表项下的缩进续行、表格（可缩进书写）、单独一行的图片 ![说明](相对路径)。
用法: python3 md_to_docx.py <输入.md> [输出.docx] [--drop-title]（图片路径相对输入文件所在目录）
  --drop-title：去掉开头的一级标题，其余标题各升一级。导入飞书时用，因为飞书文档自带标题栏，导入时用 --name 指定标题。
"""
import re, sys, copy
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

CJK_FONT, MONO_FONT = "PingFang SC", "Menlo"
INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")


def set_cjk(run, font=CJK_FONT):
    run.font.name = font
    rpr = run._element.get_or_add_rPr(); rf = rpr.find(qn("w:rFonts"))
    if rf is None: rf = OxmlElement("w:rFonts"); rpr.append(rf)
    rf.set(qn("w:eastAsia"), CJK_FONT)


def add_inline(par, text, size=None, bold=False):
    for part in INLINE.split(text):
        if not part: continue
        if part.startswith("**") and part.endswith("**"):
            r = par.add_run(part[2:-2]); r.bold = True; set_cjk(r)
        elif part.startswith("`") and part.endswith("`"):
            r = par.add_run(part[1:-1]); set_cjk(r, MONO_FONT); r.font.color.rgb = RGBColor(0x8A, 0x2B, 0x2B)
        else:
            r = par.add_run(part); set_cjk(r); r.bold = bold or None
        if size: r.font.size = Pt(size)


class Numbering:
    """每个有序列表新建一个编号实例（startOverride=1），避免 Word 把多个列表连续编号。"""
    def __init__(self, doc):
        self.doc = doc; self.numbering = doc.part.numbering_part.element
        style = doc.styles["List Number"]; numpr = style.element.pPr.find(qn("w:numPr"))
        base_num = int(numpr.find(qn("w:numId")).get(qn("w:val")))
        num = next(n for n in self.numbering.findall(qn("w:num")) if int(n.get(qn("w:numId"))) == base_num)
        self.abstract = num.find(qn("w:abstractNumId")).get(qn("w:val"))

    def new(self):
        ids = [int(n.get(qn("w:numId"))) for n in self.numbering.findall(qn("w:num"))]; nid = max(ids) + 1
        num = OxmlElement("w:num"); num.set(qn("w:numId"), str(nid))
        an = OxmlElement("w:abstractNumId"); an.set(qn("w:val"), self.abstract); num.append(an)
        ov = OxmlElement("w:lvlOverride"); ov.set(qn("w:ilvl"), "0")
        so = OxmlElement("w:startOverride"); so.set(qn("w:val"), "1"); ov.append(so); num.append(ov)
        self.numbering.append(num); return nid

    @staticmethod
    def apply(par, nid):
        ppr = par._p.get_or_add_pPr(); numpr = OxmlElement("w:numPr")
        il = OxmlElement("w:ilvl"); il.set(qn("w:val"), "0"); ni = OxmlElement("w:numId"); ni.set(qn("w:val"), str(nid))
        numpr.append(il); numpr.append(ni); ppr.append(numpr)


def add_table(doc, rows):
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    header, body = cells[0], [r for r in cells[2:]]
    ncol = len(header); t = doc.add_table(rows=1 + len(body), cols=ncol); t.style = "Table Grid"
    for i, row in enumerate([header] + body):
        for j in range(ncol):
            c = t.cell(i, j); c.text = ""; p = c.paragraphs[0]
            add_inline(p, row[j] if j < len(row) else "", size=10, bold=(i == 0))
    doc.add_paragraph()


def convert(src: Path, dst: Path, drop_title: bool = False):
    doc = Document()
    st = doc.styles["Normal"]; st.font.name = CJK_FONT; st.font.size = Pt(11)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    for name in ("Heading 1", "Heading 2", "Heading 3"):
        h = doc.styles[name]; h.font.name = CJK_FONT; h.element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT); h.font.color.rgb = RGBColor(0x1D, 0x1D, 0x1F)
    num = Numbering(doc); cur_num = None; last_list_par = False
    lines = src.read_text(encoding="utf-8").splitlines(); i = 0
    if drop_title:
        k = next(j for j, ln in enumerate(lines) if ln.strip())
        if lines[k].startswith("# "): lines = lines[:k] + lines[k + 1:]
        lines = [ln[1:] if re.match(r"^#{2,3}\s", ln) else ln for ln in lines]
    while i < len(lines):
        raw = lines[i]; s = raw.strip()
        if not s:
            i += 1; continue
        if s.startswith("|"):                                         # 表格（允许缩进）
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i].strip()); i += 1
            add_table(doc, block); continue
        m = re.match(r"^(#{1,3})\s+(.*)", raw)
        if m:
            doc.add_heading("", level=len(m.group(1))); add_inline(doc.paragraphs[-1], m.group(2))
            cur_num = None; i += 1; continue
        m = re.match(r"^!\[(.*?)\]\((.*?)\)\s*$", s)
        if m:
            img = (src.parent / m.group(2)).resolve()
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(img), width=Cm(15.5))
            if m.group(1):
                cp = doc.add_paragraph(); cp.alignment = WD_ALIGN_PARAGRAPH.CENTER; add_inline(cp, m.group(1), size=9)
            i += 1; continue
        m = re.match(r"^(\s*)-\s+(.*)", raw)
        if m:
            lvl = 1 if len(m.group(1)) >= 2 else 0
            p = doc.add_paragraph(style="List Bullet 2" if lvl else "List Bullet"); add_inline(p, m.group(2)); i += 1; continue
        m = re.match(r"^(\d+)\.\s+(.*)", raw)
        if m:
            if m.group(1) == "1" or cur_num is None: cur_num = num.new()
            p = doc.add_paragraph(style="List Number"); Numbering.apply(p, cur_num); add_inline(p, m.group(2)); i += 1; continue
        if raw.startswith("  "):                                      # 列表项下的缩进续行
            p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.9); add_inline(p, s); i += 1; continue
        cur_num = None
        p = doc.add_paragraph(); add_inline(p, s); i += 1
    doc.save(str(dst)); return dst


if __name__ == "__main__":
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    src = Path(args[0]); dst = Path(args[1]) if len(args) > 1 else src.with_suffix(".docx")
    print("saved", convert(src, dst, drop_title="--drop-title" in sys.argv))
